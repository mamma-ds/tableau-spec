"""WorkbookSpec から Excel(.xlsx) バイト列を生成する。ファイル書き込みは行わない。

reporter.py（HTML生成）とは責務が異なるため別モジュールとした。両者は独立して
WorkbookSpec からそれぞれの出力形式を組み立てるだけで、互いには依存しない。
依存関係ツリーはツリー構造で表形式に馴染まないため、Excel出力の対象外とする。
"""

from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from tableau_spec.analyzer import WorkbookSpec

_INVALID_SHEET_CHARS = set(":\\/?*[]")
_MAX_SHEET_NAME_LENGTH = 31
_MAX_COLUMN_WIDTH = 60
_WRAP_COLUMN_WIDTH = 50


def _safe_sheet_name(name: str) -> str:
    cleaned = "".join(c for c in name if c not in _INVALID_SHEET_CHARS)
    return cleaned[:_MAX_SHEET_NAME_LENGTH] or "Sheet"


def _autosize_columns(ws: Worksheet, wrap_columns: set[int] | None = None) -> None:
    wrap_columns = wrap_columns or set()
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        if col_idx in wrap_columns:
            width = _WRAP_COLUMN_WIDTH
            for cell in column_cells:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        else:
            lengths = [len(str(c.value)) for c in column_cells if c.value is not None]
            width = min(max(lengths, default=10) + 2, _MAX_COLUMN_WIDTH)
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _write_sheet(
    wb: Workbook,
    title: str,
    headers: list[str],
    rows: list[list[str]],
    wrap_columns: set[int] | None = None,
) -> None:
    ws = wb.create_sheet(title=_safe_sheet_name(title))
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    _autosize_columns(ws, wrap_columns)


def _build_field_usage_map(spec: WorkbookSpec) -> dict[str, list[str]]:
    """フィールド／計算フィールドの表示名 → 使用しているシート名一覧の対応表を作る。"""
    usage: dict[str, list[str]] = {}
    for s in spec.sheets:
        for name in list(s.used_fields) + list(s.used_calculated_fields):
            usage.setdefault(name, []).append(s.name)
    return usage


def _compute_all_used_captions(spec: WorkbookSpec) -> set[str]:
    """全シートで実際に使われている（または使われている計算フィールドの数式から
    参照される）フィールド／計算フィールドの表示名集合を返す。"""
    calc_by_caption = {f.caption: f for f in spec.calculated_fields}
    used: set[str] = set()
    for s in spec.sheets:
        used.update(s.used_fields)
        used.update(s.used_calculated_fields)
    queue = [c for c in used if c in calc_by_caption]
    while queue:
        caption = queue.pop()
        for dep in calc_by_caption[caption].depends_on:
            if dep not in used:
                used.add(dep)
                if dep in calc_by_caption:
                    queue.append(dep)
    return used


def _field_with_shelf(name: str, shelves: dict[str, list[str]]) -> str:
    roles = shelves.get(name)
    return f"{name}（{'・'.join(roles)}）" if roles else name


def _write_overview_sheet(wb: Workbook, spec: WorkbookSpec, source_name: str) -> None:
    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    rows = [
        ["ファイル名", source_name],
        ["生成日時", generated_at],
        ["データソース数", len(spec.datasources)],
        ["計算フィールド数", len(spec.calculated_fields)],
        ["パラメーター数", len(spec.parameters)],
        ["シート数", len(spec.sheets)],
        ["ダッシュボード数", len(spec.dashboards)],
    ]
    _write_sheet(wb, "概要", ["項目", "値"], rows)


def _write_datasources_sheet(wb: Workbook, spec: WorkbookSpec) -> None:
    rows = []
    for ds in spec.datasources:
        table_lines = []
        for t in ds.tables:
            table_name = t.table or t.name
            columns = "、".join(t.columns) if t.columns else "（フィールド情報なし）"
            table_lines.append(f"{table_name}: {columns}")
        tables_text = "\n".join(table_lines) if table_lines else "-"
        sql_text = "\n---\n".join(ds.custom_sql) if ds.custom_sql else "-"
        rows.append([ds.name, ds.connection_class, tables_text, sql_text])
    _write_sheet(
        wb,
        "データソース",
        ["データソース名", "接続種別", "テーブル", "カスタムSQL"],
        rows,
        wrap_columns={3, 4},
    )


def _write_joins_sheet(wb: Workbook, spec: WorkbookSpec) -> None:
    rows = [
        [ds.name, j.left_table, j.join_type, j.right_table, j.condition]
        for ds in spec.datasources
        for j in ds.joins
    ]
    _write_sheet(
        wb, "テーブル結合", ["データソース", "テーブル1", "結合種別", "テーブル2", "結合条件"], rows
    )


def _write_fields_sheet(wb: Workbook, spec: WorkbookSpec) -> None:
    usage = _build_field_usage_map(spec)
    field_names = sorted({name for s in spec.sheets for name in s.used_fields})
    rows = [[name, "、".join(usage.get(name, [])) or "-"] for name in field_names]
    _write_sheet(wb, "フィールド一覧", ["フィールド名", "使用シート"], rows)


def _write_calculated_fields_sheet(wb: Workbook, spec: WorkbookSpec) -> None:
    usage = _build_field_usage_map(spec)
    rows = [
        [
            f.caption,
            "○" if f.is_lod else "",
            f.formula,
            f.datasource,
            "、".join(usage.get(f.caption, [])) or "-",
        ]
        for f in spec.calculated_fields
    ]
    _write_sheet(
        wb,
        "計算フィールド",
        ["フィールド名", "LOD", "計算式", "データソース", "使用シート"],
        rows,
        wrap_columns={3},
    )


def _write_parameters_sheet(wb: Workbook, spec: WorkbookSpec) -> None:
    rows = [[p.caption, p.datatype, p.current_value] for p in spec.parameters]
    _write_sheet(wb, "パラメーター", ["パラメーター名", "型", "既定値"], rows)


def _write_sets_sheet(wb: Workbook, spec: WorkbookSpec) -> None:
    rows = [[s.name, s.field, s.description, s.datasource] for s in spec.sets]
    _write_sheet(wb, "セット", ["セット名", "対象フィールド", "定義", "データソース"], rows)


def _write_unused_fields_sheet(wb: Workbook, spec: WorkbookSpec) -> None:
    used = _compute_all_used_captions(spec)
    rows = [
        [caption, ds.name]
        for ds in spec.datasources
        for caption in ds.all_field_captions
        if caption not in used
    ]
    _write_sheet(wb, "未使用フィールド", ["フィールド名", "データソース"], rows)


def _write_dashboards_sheet(wb: Workbook, spec: WorkbookSpec) -> None:
    rows = [[d.name, "、".join(d.sheets) or "-"] for d in spec.dashboards]
    _write_sheet(wb, "ダッシュボード一覧", ["ダッシュボード名", "含まれるシート"], rows)


def _write_sheets_sheet(wb: Workbook, spec: WorkbookSpec) -> None:
    rows = []
    for s in spec.sheets:
        filters = "、".join(f.column for f in s.filters) or "-"
        used_fields = "、".join(_field_with_shelf(f, s.field_shelves) for f in s.used_fields) or "-"
        used_calc_fields = (
            "、".join(_field_with_shelf(f, s.field_shelves) for f in s.used_calculated_fields) or "-"
        )
        rows.append([s.name, filters, used_fields, used_calc_fields])
    _write_sheet(
        wb,
        "シート一覧",
        ["シート名", "フィルター", "使用フィールド", "使用計算フィールド"],
        rows,
        wrap_columns={3, 4},
    )


def render_excel(spec: WorkbookSpec, source_name: str) -> bytes:
    """WorkbookSpec を Excel ブック(.xlsx)のバイト列に変換する。
    HTML仕様書の「青字のタイトル」ごとの表を、依存関係ツリーを除いてシートに分けて出力する。"""
    wb = Workbook()
    wb.remove(wb.active)

    _write_overview_sheet(wb, spec, source_name)
    _write_datasources_sheet(wb, spec)
    _write_joins_sheet(wb, spec)
    _write_fields_sheet(wb, spec)
    _write_calculated_fields_sheet(wb, spec)
    _write_parameters_sheet(wb, spec)
    _write_sets_sheet(wb, spec)
    _write_unused_fields_sheet(wb, spec)
    _write_dashboards_sheet(wb, spec)
    _write_sheets_sheet(wb, spec)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
