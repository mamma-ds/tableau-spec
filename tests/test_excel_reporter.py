import io

from openpyxl import load_workbook

from tableau_spec.analyzer import (
    CalculatedField,
    DataSource,
    Dashboard,
    JoinInfo,
    Parameter,
    SetInfo,
    Sheet,
    TableInfo,
    WorkbookSpec,
)
from tableau_spec.excel_reporter import _safe_sheet_name, render_excel


def _load(data: bytes):
    return load_workbook(io.BytesIO(data))


def test_render_excel_creates_one_sheet_per_table_excluding_dependency_tree():
    spec = WorkbookSpec()

    wb = _load(render_excel(spec, "sample.twb"))

    assert wb.sheetnames == [
        "概要",
        "データソース",
        "テーブル結合",
        "フィールド一覧",
        "計算フィールド",
        "パラメーター",
        "セット",
        "未使用フィールド",
        "ダッシュボード一覧",
        "シート一覧",
    ]
    assert "依存関係ツリー" not in wb.sheetnames


def test_render_excel_overview_sheet_contains_counts():
    spec = WorkbookSpec(
        datasources=[DataSource(name="ds1", connection_class="federated")],
        sheets=[Sheet(name="シート1")],
    )

    wb = _load(render_excel(spec, "sample.twb"))
    ws = wb["概要"]

    values = {row[0]: row[1] for row in ws.iter_rows(min_row=2, values_only=True)}
    assert values["ファイル名"] == "sample.twb"
    assert values["データソース数"] == 1
    assert values["シート数"] == 1


def test_render_excel_calculated_fields_sheet_marks_lod_and_formula():
    spec = WorkbookSpec(
        calculated_fields=[
            CalculatedField(
                name="[c1]", caption="固定利益", formula="{FIXED [x]: SUM([y])}", is_lod=True, datasource="ds1"
            ),
            CalculatedField(name="[c2]", caption="通常", formula="SUM([y])", is_lod=False, datasource="ds1"),
        ],
        sheets=[Sheet(name="シート1", used_calculated_fields=["固定利益"])],
    )

    wb = _load(render_excel(spec, "sample.twb"))
    ws = wb["計算フィールド"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    lod_row = next(r for r in rows if r[0] == "固定利益")
    normal_row = next(r for r in rows if r[0] == "通常")
    assert lod_row[1] == "○"
    assert lod_row[2] == "{FIXED [x]: SUM([y])}"
    assert lod_row[4] == "シート1"
    # openpyxlは空文字セルを再読込時にNoneとして返す
    assert not normal_row[1]


def test_render_excel_sheets_sheet_annotates_field_shelves():
    spec = WorkbookSpec(
        sheets=[
            Sheet(
                name="シート1",
                used_fields=["地域"],
                used_calculated_fields=["利益率"],
                field_shelves={"地域": ["行"], "利益率": ["色", "フィルター"]},
            )
        ],
    )

    wb = _load(render_excel(spec, "sample.twb"))
    ws = wb["シート一覧"]
    row = list(ws.iter_rows(min_row=2, values_only=True))[0]

    assert row[3] == "地域（行）"
    assert row[4] == "利益率（色・フィルター）"


def test_render_excel_sheets_sheet_shows_used_datasources():
    spec = WorkbookSpec(
        sheets=[
            Sheet(name="シート1", datasources=["売上データ", "結合データ"]),
            Sheet(name="シート2"),
        ],
    )

    wb = _load(render_excel(spec, "sample.twb"))
    ws = wb["シート一覧"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    assert ws["B1"].value == "使用データソース"
    assert rows[0][1] == "売上データ、結合データ"
    assert rows[1][1] == "-"


def test_render_excel_joins_and_unused_fields_sheets():
    spec = WorkbookSpec(
        datasources=[
            DataSource(
                name="ds1",
                connection_class="federated",
                tables=[TableInfo(name="t1", table="[t1]", columns=["a", "b"])],
                all_field_captions=["a", "b", "c"],
                joins=[
                    JoinInfo(join_type="inner", left_table="[t1]", right_table="[t2]", condition="[t1].[id] = [t2].[id]")
                ],
            )
        ],
        sheets=[Sheet(name="シート1", used_fields=["a"])],
    )

    wb = _load(render_excel(spec, "sample.twb"))

    joins_rows = list(wb["テーブル結合"].iter_rows(min_row=2, values_only=True))
    assert joins_rows == [("ds1", "[t1]", "inner", "[t2]", "[t1].[id] = [t2].[id]")]

    unused_rows = list(wb["未使用フィールド"].iter_rows(min_row=2, values_only=True))
    assert unused_rows == [("b", "ds1"), ("c", "ds1")]


def test_render_excel_sets_and_dashboards_sheets():
    spec = WorkbookSpec(
        sets=[SetInfo(name="上位顧客", field="顧客名", description="上位 5 件", datasource="ds1")],
        dashboards=[Dashboard(name="ダッシュボード1", sheets=["シートA", "シートB"])],
    )

    wb = _load(render_excel(spec, "sample.twb"))

    sets_rows = list(wb["セット"].iter_rows(min_row=2, values_only=True))
    assert sets_rows == [("上位顧客", "顧客名", "上位 5 件", "ds1")]

    dash_rows = list(wb["ダッシュボード一覧"].iter_rows(min_row=2, values_only=True))
    assert dash_rows == [("ダッシュボード1", "シートA、シートB")]


def test_render_excel_header_row_is_bold():
    spec = WorkbookSpec(sheets=[Sheet(name="シート1")])

    wb = _load(render_excel(spec, "sample.twb"))
    ws = wb["シート一覧"]

    assert ws["A1"].font.bold is True


def test_safe_sheet_name_strips_invalid_characters_and_truncates():
    assert _safe_sheet_name("a:b/c*d") == "abcd"
    assert _safe_sheet_name("あ" * 40) == "あ" * 31
    assert _safe_sheet_name("") == "Sheet"
